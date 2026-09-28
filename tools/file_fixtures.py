"""Public synthetic Korean PDF/HWPX fixtures, never user documents."""
from io import BytesIO
from pathlib import Path
import zipfile

from PIL import Image, ImageDraw, ImageFont


def create_fixtures(directory):
    from reportlab.pdfgen import canvas
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.lib.utils import ImageReader
    directory = Path(directory)
    directory.mkdir(parents=True, exist_ok=True)
    font_path = Path('C:/Windows/Fonts/malgun.ttf')
    pdfmetrics.registerFont(TTFont('SsoklyFixture', str(font_path)))
    scan = Image.new('RGB', (1400, 800), 'white')
    draw = ImageDraw.Draw(scan)
    font = ImageFont.truetype(str(font_path), 38)
    lines = ['합성 독서교실 안내', '신청 마감: 2026년 10월 2일 오후 4시',
             '행사: 2026년 10월 16일 오후 2시', '대상: 희망하는 3학년 학생 / 참가비: 무료']
    for index, line in enumerate(lines):
        draw.text((50, 60 + index * 105), line, fill='black', font=font)
    scan.save(directory / 'scan-source.png')
    pdf_path = directory / 'synthetic-mixed.pdf'
    pdf = canvas.Canvas(str(pdf_path), pagesize=(595, 842))
    pdf.setFont('SsoklyFixture', 18)
    pdf.drawString(42, 785, '합성 학교 행사 일정')
    pdf.setFont('SsoklyFixture', 12)
    for y, values in ((720, ['프로그램', '신청 마감', '행사일']),
                       (682, ['독서교실', '2026. 10. 2. 16:00', '2026. 10. 16. 14:00']),
                       (644, ['과학교실', '2026. 10. 6. 15:00', '2026. 10. 20. 10:00'])):
        for x, value in zip((42, 185, 380), values):
            pdf.drawString(x, y, value)
    pdf.showPage()
    pdf.drawImage(ImageReader(scan), 30, 280, width=535, height=535 * 800 / 1400)
    pdf.showPage()
    pdf.save()
    def paragraph(text):
        return '<p><run><t>' + text + '</t></run></p>'
    body = paragraph('합성 학교 행사 일정')
    body += '<tbl rowCnt="2" colCnt="3">'
    for row, values in enumerate((['프로그램', '신청 마감', '행사일'], ['독서교실', '2026. 10. 2. 16:00', '2026. 10. 16. 14:00'])):
        body += '<tr>'
        for col, value in enumerate(values):
            body += f'<tc><cellAddr rowAddr="{row}" colAddr="{col}"/><cellSpan rowSpan="1" colSpan="1"/><subList>{paragraph(value)}</subList></tc>'
        body += '</tr>'
    body += '</tbl><p><run><pic><img binaryItemIDRef="image1"/></pic></run></p>'
    image_bytes = BytesIO()
    scan.save(image_bytes, format='PNG')
    hwpx_path = directory / 'synthetic-table-image.hwpx'
    with zipfile.ZipFile(hwpx_path, 'w') as archive:
        archive.writestr('Contents/section0.xml', '<sec>' + body + '</sec>')
        archive.writestr('Contents/content.hpf', '<package><manifest><item id="image1" href="BinData/image1.png"/></manifest></package>')
        archive.writestr('BinData/image1.png', image_bytes.getvalue())
    scan.close()
    return pdf_path, hwpx_path


def create_supported_fixtures(directory):
    from docx import Document
    from pptx import Presentation
    from pptx.util import Inches
    from openpyxl import Workbook
    from services.document_service import IMAGE_EXTENSIONS, TEXT_EXTENSIONS
    directory = Path(directory)
    pdf, hwpx = create_fixtures(directory)
    result = {'.pdf': pdf, '.hwpx': hwpx}
    content = '합성 문서\n신청 마감\t2026. 10. 2. 16:00'
    for extension in TEXT_EXTENSIONS:
        path = directory / ('synthetic' + extension)
        if extension == '.json':
            import json
            value = json.dumps({'title': '합성 문서', 'deadline': '2026. 10. 2. 16:00'}, ensure_ascii=False)
        elif extension == '.xml':
            value = '<document><title>합성 문서</title><deadline>2026. 10. 2. 16:00</deadline></document>'
        elif extension == '.html':
            value = '<!doctype html><html><body><h1>합성 문서</h1><p>2026. 10. 2. 16:00</p></body></html>'
        else:
            value = content
        path.write_text(value, encoding='utf-8')
        result[extension] = path
    for extension in IMAGE_EXTENSIONS:
        path = directory / ('synthetic' + extension)
        with Image.open(directory / 'scan-source.png') as image:
            image.save(path)
        result[extension] = path
    document = Document()
    document.add_paragraph('합성 문서')
    table = document.add_table(rows=2, cols=2)
    for row, values in zip(table.rows, [('항목', '기한'), ('신청 마감', '2026. 10. 2. 16:00')]):
        for cell, value in zip(row.cells, values):
            cell.text = value
    path = directory / 'synthetic.docx'
    document.save(path)
    result['.docx'] = path
    presentation = Presentation()
    for number in (1, 2):
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        slide.shapes.add_textbox(Inches(1), Inches(1), Inches(8), Inches(2)).text_frame.text = content + f'\n슬라이드 {number}'
    path = directory / 'synthetic.pptx'
    presentation.save(path)
    result['.pptx'] = path
    workbook = Workbook()
    workbook.active.title = '일정'
    workbook.active.append(['합성 문서', '기한'])
    workbook.active.append(['신청 마감', '2026. 10. 2. 16:00'])
    workbook.create_sheet('추가').append(['합성 문서', '두 번째 시트'])
    path = directory / 'synthetic.xlsx'
    workbook.save(path)
    workbook.close()
    result['.xlsx'] = path
    rtf = ''.join('\\u' + str(ord(c) if ord(c) < 32768 else ord(c) - 65536) + '?' if ord(c) > 127 else c for c in '합성 문서 2026. 10. 2. 16:00')
    path = directory / 'synthetic.rtf'
    path.write_text('{\\rtf1\\ansi\\uc1 ' + rtf + '}', encoding='ascii')
    result['.rtf'] = path
    namespaces = ('xmlns:office="urn:oasis:names:tc:opendocument:xmlns:office:1.0" '
                  'xmlns:text="urn:oasis:names:tc:opendocument:xmlns:text:1.0" '
                  'xmlns:table="urn:oasis:names:tc:opendocument:xmlns:table:1.0" '
                  'xmlns:draw="urn:oasis:names:tc:opendocument:xmlns:drawing:1.0"')
    paragraph = '<text:p>합성 문서 2026. 10. 2. 16:00</text:p>'
    bodies = {'.odt': '<office:text>' + paragraph + '</office:text>',
              '.odp': '<office:presentation><draw:page>' + paragraph + '</draw:page></office:presentation>',
              '.ods': '<office:spreadsheet><table:table table:name="일정"><table:table-row><table:table-cell>' + paragraph + '</table:table-cell></table:table-row></table:table></office:spreadsheet>'}
    for extension, body in bodies.items():
        path = directory / ('synthetic' + extension)
        with zipfile.ZipFile(path, 'w') as archive:
            archive.writestr('content.xml', '<office:document-content ' + namespaces + '><office:body>' + body + '</office:body></office:document-content>')
        result[extension] = path
    return result
