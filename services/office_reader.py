"""Local Office/ODF text readers. These do not reproduce printed layouts."""
from datetime import date, datetime
import zipfile
from xml.etree import ElementTree as ET


def _cell(text):
    return str(text).replace('\t', ' | ').replace('\n', ' / ')


def _table(rows):
    return '\n'.join('\t'.join(_cell(value) for value in row) for row in rows)


def _entry(path, text, suffix=''):
    return {'text': text, 'image': None, 'name': path.name + suffix}


def _local(tag):
    return tag.rsplit('}', 1)[-1]


def _attribute(node, name, default=None):
    return next((v for k, v in node.attrib.items() if _local(k) == name), default)


def _odf_text(node):
    text = node.text or ''
    for child in node:
        kind = _local(child.tag)
        if kind == 's':
            text += ' ' * min(1000, int(_attribute(child, 'c', '1')))
        elif kind in ('tab', 'line-break'):
            text += '\t' if kind == 'tab' else '\n'
        else:
            text += _odf_text(child)
        text += child.tail or ''
    return text


def _odf_blocks(node):
    kind = _local(node.tag)
    if kind in ('p', 'h'):
        return [_odf_text(node)]
    if kind == 'table':
        rows = []
        count = 0
        for row in node.iter():
            if _local(row.tag) != 'table-row':
                continue
            values = []
            repeat = int(_attribute(row, 'number-rows-repeated', '1'))
            for cell in row:
                if _local(cell.tag) not in ('table-cell', 'covered-table-cell'):
                    continue
                times = int(_attribute(cell, 'number-columns-repeated', '1'))
                if not 1 <= times <= 1024:
                    raise ValueError('ODF 반복 셀 범위가 너무 큽니다. 필요한 범위만 별도 저장하세요.')
                value = ' / '.join(_odf_text(p) for p in cell if _local(p.tag) in ('p', 'h'))
                if not value:
                    value = _attribute(cell, 'date-value') or _attribute(cell, 'value') or ''
                if _local(cell.tag) == 'covered-table-cell':
                    value = '↳'
                values.extend([value] * times)
            if not 1 <= repeat <= 10000 or len(values) > 1024:
                raise ValueError('ODF 표 범위가 너무 큽니다.')
            count += repeat * len(values)
            if count > 100000:
                raise ValueError('ODF 표는 10만 셀까지 읽을 수 있습니다.')
            rows.extend([values] * repeat)
        return [_table(rows)]
    return [text for child in node for text in _odf_blocks(child)]


def read_office(path):
    suffix = path.suffix.lower()
    if suffix != '.rtf':
        with zipfile.ZipFile(path) as archive:
            if sum(item.file_size for item in archive.infolist()) > 128 * 1024 * 1024:
                raise ValueError('압축 해제 크기가 128MB를 초과합니다.')
    if suffix == '.docx':
        from docx import Document
        from docx.table import Table
        document = Document(path)
        def blocks(container):
            result = []
            for block in container.iter_inner_content():
                if isinstance(block, Table):
                    result.append(_table([[' / '.join(blocks(cell)) for cell in row.cells] for row in block.rows]))
                else:
                    result.append(block.text)
            return result
        return [_entry(path, '\n'.join(blocks(document)))]
    if suffix == '.pptx':
        from pptx import Presentation
        def shapes_text(shapes):
            result = []
            for shape in shapes:
                if hasattr(shape, 'shapes'):
                    result.extend(shapes_text(shape.shapes))
                elif shape.has_table:
                    result.append(_table([[cell.text for cell in row.cells] for row in shape.table.rows]))
                elif shape.has_text_frame:
                    result.append(shape.text)
            return result
        return [_entry(path, '\n'.join(shapes_text(slide.shapes)), f' · 슬라이드 {i + 1}')
                for i, slide in enumerate(Presentation(path).slides)]
    if suffix == '.xlsx':
        from openpyxl import load_workbook
        values = load_workbook(path, read_only=True, data_only=True, keep_links=False)
        formulas = None
        try:
            formulas = load_workbook(path, read_only=True, data_only=False, keep_links=False)
            result = []
            for sheet in values:
                if sheet.sheet_state != 'visible':
                    continue
                if (sheet.max_row or 0) * (sheet.max_column or 0) > 100000:
                    raise ValueError('시트는 10만 셀까지 읽을 수 있습니다. 필요한 범위만 별도 저장하세요.')
                rows = []
                cell_count = 0
                for vr, fr in zip(sheet.iter_rows(), formulas[sheet.title].iter_rows()):
                    cell_count += len(vr)
                    if cell_count > 100000:
                        raise ValueError('시트는 10만 셀까지 읽을 수 있습니다.')
                    row = []
                    for cell, formula in zip(vr, fr):
                        value = cell.value
                        if value is None and formula.data_type == 'f':
                            value = '[계산 결과 없음: ' + str(formula.value) + ']'
                        elif isinstance(value, (datetime, date)):
                            value = value.isoformat(sep=' ') if isinstance(value, datetime) else value.isoformat()
                        elif isinstance(value, (int, float)) and '%' in cell.number_format:
                            value = f'{value * 100:g}%'
                        row.append('' if value is None else str(value))
                    rows.append(row)
                result.append(_entry(path, _table(rows), ' · 시트 ' + sheet.title))
            return result
        finally:
            values.close()
            if formulas is not None:
                formulas.close()
    if suffix == '.rtf':
        from striprtf.striprtf import rtf_to_text
        raw = path.read_bytes().decode('latin1')
        if not raw.lstrip().startswith('{\\rtf'):
            raise ValueError('올바른 RTF 파일이 아닙니다.')
        return [_entry(path, rtf_to_text(raw))]
    with zipfile.ZipFile(path) as archive:
        root = ET.fromstring(archive.read('content.xml'))
    body = next(node for node in root if _local(node.tag) == 'body')
    if suffix == '.odp':
        return [_entry(path, '\n'.join(_odf_blocks(node)), f' · 슬라이드 {i + 1}')
                for i, node in enumerate(n for n in body.iter() if _local(n.tag) == 'page')]
    if suffix == '.ods':
        return [_entry(path, '\n'.join(_odf_blocks(node)), ' · 시트 ' + _attribute(node, 'name', str(i + 1)))
                for i, node in enumerate(n for n in body.iter() if _local(n.tag) == 'table')]
    return [_entry(path, '\n'.join(_odf_blocks(body)))]
