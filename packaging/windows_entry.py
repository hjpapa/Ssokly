"""Frozen desktop entry; optional isolated offline diagnostics for packaging."""
import json
import os
from contextlib import contextmanager
from pathlib import Path
import sys
import tempfile


@contextmanager
def isolated_user_paths(root):
    """Exercise normal startup without touching the user's library or desktop."""
    from services import capture_location
    previous_local = os.environ.get('LOCALAPPDATA')
    previous_desktop = capture_location.desktop_directory
    try:
        os.environ['LOCALAPPDATA'] = str(root / 'localappdata')
        capture_location.desktop_directory = lambda: root / 'desktop'
        yield
    finally:
        capture_location.desktop_directory = previous_desktop
        if previous_local is None:
            os.environ.pop('LOCALAPPDATA', None)
        else:
            os.environ['LOCALAPPDATA'] = previous_local


def self_test(report_path):
    import tkinter as tk
    import httpx
    import openai
    import mss
    import docx
    import pptx
    import openpyxl
    import pypdfium2 as pdfium
    from PIL import Image
    from striprtf.striprtf import rtf_to_text
    from main import _enable_windows_dpi_awareness
    from ui.capture_desk import CaptureDeskApp
    from services.document_library import DocumentLibrary
    from services.capture_store import CaptureStore
    from services.ai_relay import server_url
    from services.source_review import source_table_blocks
    from services.work_image import image_request, validate_png
    from ui.work_image import WorkImageWindow
    import io
    import time
    _enable_windows_dpi_awareness()
    with tempfile.TemporaryDirectory(prefix='ssokly-package-test-') as temporary:
        root = Path(temporary)
        # Load the native PDFium library and render a page, not just import it.
        pdf = pdfium.PdfDocument.new()
        page = pdf.new_page(200, 200)
        bitmap = page.render()
        assert bitmap.to_pil().size == (200, 200)
        bitmap.close()
        page.close()
        pdf.close()
        docx.Document().save(root / 'sample.docx')
        docx.Document(root / 'sample.docx')
        pptx.Presentation().save(root / 'sample.pptx')
        pptx.Presentation(root / 'sample.pptx')
        workbook = openpyxl.Workbook()
        workbook.save(root / 'sample.xlsx')
        workbook.close()
        openpyxl.load_workbook(root / 'sample.xlsx').close()
        assert rtf_to_text(r'{\rtf1 synthetic}') == 'synthetic'
        with isolated_user_paths(root):
            app = CaptureDeskApp()
        try:
            app.withdraw()
            initial_documents = len(app.library.list_documents())
            initial_trash_documents = len(app.library.list_documents(trashed=True))
            initial_captures = len(app.library.capture_store.list_recent())
            initial_trash_captures = len(app.library.capture_store.search(trashed=True))
            assert (initial_documents, initial_trash_documents,
                    initial_captures, initial_trash_captures) == (0, 0, 0, 0)
            assert app.document is None and app.page is None
            from services.image_clipboard import copy_image
            assert callable(copy_image)
            assert str(app.image_copy_button.cget('state')) == 'disabled'
            assert app.capture_copy_button.cget('text') == '캡처 후 복사'
            assert not app.document_tree.get_children()
            manual = app.show_help()
            app.update_idletasks()
            assert 'Ssokly' in manual.text.get('1.0', 'end')
            assert app.show_help() is manual
            manual.close()
            assert app.library.app_data_dir == (root / 'localappdata' / 'Ssokly').resolve()
            assert app.library.capture_store.directory == (root / 'desktop' / 'ssokly').resolve()
            app.accept_capture(Image.new('RGB', (100, 100), 'white'), auto_read=False)
            app.update_idletasks()
            assert len(app.library.list_documents()) == 1
            assert str(app.image_copy_button.cget('state')) == 'normal'
            assert len(source_table_blocks('A | B\nC | D')) == 1
            sample = io.BytesIO()
            Image.new('RGB', (100, 150), 'white').save(sample, 'PNG')
            assert image_request('합성 업무')['model'] == 'gpt-image-2.5-flare'
            dialog = WorkImageWindow(app, '합성 업무 · 희망자만')
            dialog.started = time.monotonic()
            dialog.results.put((validate_png(sample.getvalue()), '합성 업무 · 희망자만', None))
            dialog.poll()
            app.update_idletasks()
            assert dialog.data == sample.getvalue()
            assert dialog.view._image.size == (100, 150)
            dialog.close()
        finally:
            for callback in app.tk.splitlist(app.tk.call('after', 'info')):
                app.after_cancel(callback)
            app.update_idletasks()
            app.destroy()
        reopened = DocumentLibrary(root / 'localappdata' / 'Ssokly',
                                   capture_store=CaptureStore(root / 'desktop' / 'ssokly'))
        assert len(reopened.list_documents()) == 1
    Path(report_path).write_text(json.dumps({
        'passed': True, 'frozen': bool(getattr(sys, 'frozen', False)),
        'python': sys.version.split()[0], 'relay': server_url(),
        'checks': ['Tk', 'PDFium render', 'DOCX', 'PPTX', 'XLSX', 'RTF',
                   'empty first launch', 'offline user manual', 'capture persistence', 'pipe table',
                   'OpenAI/httpx imports', 'work image preview', 'capture image copy controls'],
        'initial_documents': initial_documents,
        'initial_trash_documents': initial_trash_documents,
        'initial_captures': initial_captures,
        'initial_trash_captures': initial_trash_captures,
        'api_calls': 0,
    }, indent=2), encoding='utf-8')


if __name__ == '__main__':
    if len(sys.argv) == 3 and sys.argv[1] == '--self-test-report':
        try:
            self_test(sys.argv[2])
        except Exception as error:
            Path(sys.argv[2]).write_text(json.dumps({'passed': False, 'error_type': type(error).__name__}), encoding='utf-8')
            raise
    else:
        from main import main
        main()
